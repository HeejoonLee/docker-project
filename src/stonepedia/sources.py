"""Turning a submitted Source into a folder the agent can read and Articles can cite.

A Proposal's Sources and its Request are each prepared into their own folder, staged under
<data>/sources/<proposal id>-<n>/ and committed to the Bundle as sources/<proposal id>-<n>/ once an
accepted Article cites it. (Proposals from before 0.4 used sources/<proposal id>/.) A folder holds:

- the original (`original.<ext>`, `page.html` for links, `request.md` for Requests)
- `text.md`: extracted text, with `## Page N` / `## Slide N` / `## Sheet: X` markers
- `meta.yaml`: kind, title, filename/url, who submitted it and when, sha256
"""

import hashlib
import re
from datetime import UTC, datetime
from pathlib import Path

import httpx
import yaml
from markdownify import markdownify

from stonepedia.models import SourceInfo, SourceInput

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".heic"}
TEXT_EXTS = {".txt", ".md", ".markdown", ".csv", ".log", ".json", ".yaml", ".yml", ".xml"}
SUPPORTED_EXTS = IMAGE_EXTS | TEXT_EXTS | {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm"}


class SourceError(ValueError):
    """The Source cannot be used (unsupported, empty, unreachable...)."""


def prepare(source: SourceInput, upload: Path | None, dest: Path, submitted_by: str, title: str,
            max_bytes: int, verify_tls: bool) -> SourceInfo:
    dest.mkdir(parents=True, exist_ok=True)
    warnings: list[str] = []
    pages: int | None = None
    url = ""
    filename = ""
    title = title or source.title.strip()
    if source.kind in ("text", "request", "code", "audit"):
        original = "text.md" if source.kind == "text" else "request.md"
        body = source.text.strip()
        if not body:
            raise SourceError("The text is empty.")
        if source.kind in ("request", "code", "audit"):
            stamp = datetime.now(UTC).strftime("%Y-%m-%d")
            (dest / original).write_text(f"> Request by {submitted_by} on {stamp}\n\n{body}\n")
        text = body
        media = "code" if source.kind == "code" else "text"
        if source.kind == "text" and not title:
            title = body.splitlines()[0][:80]
    elif source.kind == "file":
        if upload is None or not upload.exists():
            raise SourceError("The uploaded file is missing.")
        filename = source.filename or upload.name
        ext = Path(filename).suffix.lower()
        if ext not in SUPPORTED_EXTS:
            raise SourceError(f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTS))}")
        original = f"original{ext}"
        upload.replace(dest / original)
        text, pages, media, warnings = extract(dest / original)
    elif source.kind == "link":
        url = source.url.strip()
        original, text, pages, media, warnings, page_title = fetch(url, dest, max_bytes, verify_tls)
        title = title or page_title
    else:
        raise SourceError(f"Unknown source kind {source.kind}")

    if source.kind != "text" or not (dest / "text.md").exists():
        (dest / "text.md").write_text(text + ("\n" if text and not text.endswith("\n") else ""))
    sha = hashlib.sha256((dest / original).read_bytes()).hexdigest()
    info = SourceInfo(
        id=dest.name,
        kind=source.kind,
        title=title or filename or url or dest.name,
        original=original,
        filename=filename,
        url=url,
        media=media,
        pages=pages,
        text_chars=len(text),
        submitted_by=submitted_by,
        submitted_at=datetime.now(UTC).isoformat(timespec="seconds"),
        sha256=sha,
        warnings=warnings,
    )
    meta = info.model_dump(exclude={"warnings"}, exclude_defaults=False)
    (dest / "meta.yaml").write_text(yaml.safe_dump(meta, sort_keys=False, allow_unicode=True))
    return info


def extract(path: Path) -> tuple[str, int | None, str, list[str]]:
    """Returns (markdown text, page/slide count, media label, warnings)."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".docx":
        return _docx(path), None, "docx", []
    if ext == ".pptx":
        text, slides = _pptx(path)
        return text, slides, "pptx", []
    if ext == ".xlsx":
        return _xlsx(path), None, "xlsx", []
    if ext in (".html", ".htm"):
        _, text = html_to_markdown(path.read_text(errors="replace"))
        return text, None, "html", []
    if ext in IMAGE_EXTS:
        return "", None, "image", ["Image: there is no extracted text; the agent reads the image itself."]
    return path.read_text(errors="replace"), None, "text", []


def fetch(url: str, dest: Path, max_bytes: int, verify_tls: bool):
    if not re.match(r"^https?://", url):
        raise SourceError("Links must start with http:// or https://")
    try:
        with httpx.Client(follow_redirects=True, timeout=30, verify=verify_tls) as client:
            with client.stream("GET", url, headers={"User-Agent": "Stonepedia"}) as resp:
                if resp.status_code >= 400:
                    raise SourceError(f"The link returned HTTP {resp.status_code}.")
                content_type = resp.headers.get("content-type", "").split(";")[0].strip().lower()
                data = b""
                for chunk in resp.iter_bytes():
                    data += chunk
                    if len(data) > max_bytes:
                        raise SourceError(f"The linked document is larger than {max_bytes // 2**20} MB.")
                final_url = str(resp.url)
    except httpx.HTTPError as err:
        raise SourceError(f"Could not fetch the link: {err}") from err

    if "login" in final_url.lower() or "sso" in final_url.lower():
        warning = [f"The link redirected to {final_url}; it may need a login. If the text looks wrong, save the page as PDF and upload it."]
    else:
        warning = []
    if content_type == "application/pdf" or final_url.lower().endswith(".pdf"):
        (dest / "original.pdf").write_bytes(data)
        text, pages, media, warnings = _pdf(dest / "original.pdf")
        return "original.pdf", text, pages, media, warning + warnings, ""
    if content_type.startswith("text/html") or not content_type:
        html = data.decode(_charset(resp.headers.get("content-type", "")), errors="replace")
        (dest / "page.html").write_text(html)
        title, text = html_to_markdown(html)
        if len(text.strip()) < 200:
            warning.append("Very little text was found on the page; it may need a login or JavaScript.")
        return "page.html", f"Source: {final_url}\n\n{text}", None, "html", warning, title
    if content_type.startswith("text/"):
        (dest / "original.txt").write_bytes(data)
        return "original.txt", data.decode(errors="replace"), None, "text", warning, ""
    raise SourceError(f"Unsupported content type '{content_type}' at the link.")


def html_to_markdown(html: str) -> tuple[str, str]:
    title_match = re.search(r"<title[^>]*>(.*?)</title>", html, re.I | re.S)
    title = re.sub(r"\s+", " ", title_match.group(1)).strip() if title_match else ""
    body = re.sub(r"<(script|style|noscript|nav|footer|header|svg)\b.*?</\1>", "", html, flags=re.I | re.S)
    text = markdownify(body, heading_style="ATX", strip=["img"])
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return title, text


def _charset(content_type: str) -> str:
    match = re.search(r"charset=([\w-]+)", content_type, re.I)
    return match.group(1) if match else "utf-8"


def _pdf(path: Path) -> tuple[str, int, str, list[str]]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(path)
        parts = []
        for i, page in enumerate(reader.pages, 1):
            parts.append(f"## Page {i}\n\n{(page.extract_text() or '').strip()}")
    except Exception as err:  # pypdf raises many types for broken files
        raise SourceError(f"Could not read the PDF: {err}") from err
    text = "\n\n".join(parts)
    warnings = []
    if sum(len(p) for p in parts) < 50 * max(1, len(parts)):
        warnings.append("The PDF has little or no text (scanned?); the agent will read the pages directly.")
    return text, len(parts), "pdf", warnings


def _docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    out: list[str] = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            from docx.text.paragraph import Paragraph

            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "").lower()
            level = re.search(r"heading (\d)", style)
            if level:
                out.append("#" * min(int(level.group(1)) + 1, 6) + " " + text)
            elif "title" in style:
                out.append("# " + text)
            elif "list" in style:
                out.append("- " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            from docx.table import Table

            rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in Table(child, doc).rows]
            out.append(_table(rows))
    return "\n\n".join(out)


def _pptx(path: Path) -> tuple[str, int]:
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        parts = [f"## Slide {i}" + (f": {title}" if title else "")]
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        parts.append("  " * para.level + "- " + text)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in shape.table.rows]
                parts.append(_table(rows))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notes: {notes}")
        out.append("\n".join(parts))
    return "\n\n".join(out), len(prs.slides)


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).replace("\n", " ").strip() for v in row]
            if any(cells):
                rows.append(cells)
        if rows:
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            out.append(f"## Sheet: {ws.title}\n\n{_table(rows)}")
    return "\n\n".join(out)


def _table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [[c.replace("|", "\\|") for c in r] + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)
