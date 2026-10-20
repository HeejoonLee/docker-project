import http.server
import threading
from pathlib import Path

import pytest
import yaml

from stonepedia import sources
from stonepedia.models import SourceInput

def build_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page."""
    n = len(pages)
    objs = {1: "<</Type/Catalog/Pages 2 0 R>>", 3: "<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>"}
    kids = []
    for i, text in enumerate(pages):
        page_id, content_id = 4 + 2 * i, 5 + 2 * i
        stream = f"BT /F1 18 Tf 20 100 Td ({text}) Tj ET"
        objs[page_id] = (f"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 144]/Contents {content_id} 0 R"
                         "/Resources<</Font<</F1 3 0 R>>>>>>")
        objs[content_id] = f"<</Length {len(stream)}>>stream\n{stream}\nendstream"
        kids.append(f"{page_id} 0 R")
    objs[2] = f"<</Type/Pages/Kids[{' '.join(kids)}]/Count {n}>>"
    out = b"%PDF-1.4\n"
    offsets = {}
    for num in sorted(objs):
        offsets[num] = len(out)
        out += f"{num} 0 obj{objs[num]}endobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for num in sorted(objs):
        out += f"{offsets[num]:010d} 00000 n \n".encode()
    out += f"trailer<</Size {len(objs) + 1}/Root 1 0 R>>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_pdf(path: Path) -> Path:
    path.write_bytes(build_pdf(["Boot from LU 1 first", "Then enter Factory"]))
    return path


def make_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("부트 절차", level=1)
    doc.add_paragraph("Factory 모드로 진입한다.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "LU", "용도"
    table.cell(1, 0).text, table.cell(1, 1).text = "1", "Boot"
    doc.save(path)
    return path


def make_pptx(path: Path) -> Path:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "FFU 흐름"
    slide.placeholders[1].text_frame.text = "Factory → DM → FFU"
    prs.save(path)
    return path


def make_xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "WBS"
    ws.append(["ID", "작업"])
    ws.append(["BL-01", "Baseline 선택"])
    wb.save(path)
    return path


def test_pdf_pages(tmp_path):
    text, pages, media, warnings = sources.extract(make_pdf(tmp_path / "a.pdf"))
    assert (pages, media) == (2, "pdf")
    assert "## Page 1\n\nBoot from LU 1 first" in text and "## Page 2\n\nThen enter Factory" in text


def test_docx_headings_and_tables(tmp_path):
    text, _, media, _ = sources.extract(make_docx(tmp_path / "a.docx"))
    assert media == "docx"
    assert "## 부트 절차" in text and "Factory 모드로 진입한다." in text
    assert "| LU | 용도 |" in text and "| 1 | Boot |" in text


def test_pptx_slides(tmp_path):
    text, slides, _, _ = sources.extract(make_pptx(tmp_path / "a.pptx"))
    assert slides == 1 and "## Slide 1: FFU 흐름" in text and "- Factory → DM → FFU" in text


def test_xlsx_sheets(tmp_path):
    text, _, _, _ = sources.extract(make_xlsx(tmp_path / "a.xlsx"))
    assert "## Sheet: WBS" in text and "| BL-01 | Baseline 선택 |" in text


def test_html_drops_scripts_and_keeps_title(tmp_path):
    title, text = sources.html_to_markdown(
        "<html><head><title>LU 가이드</title><script>alert(1)</script></head>"
        "<body><nav>menu</nav><h1>LU</h1><p>Boot LU는 <b>1</b>이다.</p></body></html>"
    )
    assert title == "LU 가이드" and "# LU" in text and "Boot LU는 **1**이다." in text
    assert "alert" not in text and "menu" not in text


def test_prepare_file_writes_folder(tmp_path):
    upload = make_docx(tmp_path / "절차서.docx")
    info = sources.prepare(
        SourceInput(kind="file", filename="절차서.docx"), upload, tmp_path / "src1", "kim", "", 2**20, False
    )
    folder = tmp_path / "src1"
    assert info.original == "original.docx" and info.title == "절차서.docx" and info.media == "docx"
    assert sorted(p.name for p in folder.iterdir()) == ["meta.yaml", "original.docx", "text.md"]
    meta = yaml.safe_load((folder / "meta.yaml").read_text())
    assert meta["filename"] == "절차서.docx" and meta["submitted_by"] == "kim" and len(meta["sha256"]) == 64


def test_prepare_image_and_request(tmp_path):
    img = tmp_path / "note.jpg"
    img.write_bytes(b"\xff\xd8\xff\xe0 not really a jpeg")
    info = sources.prepare(SourceInput(kind="file", filename="note.jpg"), img, tmp_path / "i", "kim", "", 2**20, False)
    assert info.media == "image" and info.warnings and (tmp_path / "i" / "original.jpg").exists()

    info = sources.prepare(SourceInput(kind="request", text="FFU 후 재부팅 필요"), None, tmp_path / "r", "lee", "", 2**20, False)
    assert info.original == "request.md"
    assert (tmp_path / "r" / "request.md").read_text().startswith("> Request by lee on ")


def test_unsupported_and_empty_sources(tmp_path):
    bad = tmp_path / "a.exe"
    bad.write_bytes(b"MZ")
    with pytest.raises(sources.SourceError, match="Unsupported"):
        sources.prepare(SourceInput(kind="file", filename="a.exe"), bad, tmp_path / "x", "kim", "", 2**20, False)
    with pytest.raises(sources.SourceError, match="empty"):
        sources.prepare(SourceInput(kind="text", text="  "), None, tmp_path / "y", "kim", "", 2**20, False)


@pytest.fixture
def web_server(tmp_path):
    root = tmp_path / "site"
    root.mkdir()
    (root / "page.html").write_text(
        "<html><head><meta charset='utf-8'><title>UFS 노트</title></head><body><h1>UFS</h1>"
        + "<p>Boot LU 설정은 provisioning XML로 정의한다.</p>" * 10 + "</body></html>"
    )
    make_pdf(root / "spec.pdf")
    handler = lambda *a, **kw: http.server.SimpleHTTPRequestHandler(*a, directory=str(root), **kw)  # noqa: E731
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{server.server_port}"
    server.shutdown()


def test_link_html_and_pdf(tmp_path, web_server):
    info = sources.prepare(SourceInput(kind="link", url=f"{web_server}/page.html"), None, tmp_path / "l1", "kim", "", 2**20, False)
    assert info.title == "UFS 노트" and info.original == "page.html" and info.media == "html"
    assert "provisioning XML" in (tmp_path / "l1" / "text.md").read_text()

    info = sources.prepare(SourceInput(kind="link", url=f"{web_server}/spec.pdf"), None, tmp_path / "l2", "kim", "Spec", 2**20, False)
    assert info.original == "original.pdf" and info.pages == 2 and info.title == "Spec"


def test_link_errors(tmp_path, web_server):
    with pytest.raises(sources.SourceError, match="404"):
        sources.prepare(SourceInput(kind="link", url=f"{web_server}/missing"), None, tmp_path / "e", "kim", "", 2**20, False)
    with pytest.raises(sources.SourceError, match="larger than"):
        sources.prepare(SourceInput(kind="link", url=f"{web_server}/spec.pdf"), None, tmp_path / "f", "kim", "", 100, False)
